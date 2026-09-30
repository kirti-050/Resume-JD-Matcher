import os 
import json
from datetime import datetime 

import streamlit as st
from pydantic import BaseModel
from groq import Groq 
from pypdf import PdfReader
from docx import Document
from textwrap import dedent

from dotenv import load_dotenv
load_dotenv()

# ============================================================
# SETUP
# ============================================================

def get_api_key():
    try:
        if "GROQ_API_KEY" in st.secrets:
            return st.secrets["GROQ_API_KEY"]
    except Exception:
        pass 
    return os.getenv("GROQ_API_KEY")


api_key = get_api_key() 
if not api_key:
    st.error("Groq API Key not found. Add it to streamlit secrets or a local .env file.")
    st.stop()


client = Groq(api_key = api_key)
model = "openai/gpt-oss-120b"

MAX_CHARS = 6000
COOLDOWN_SECONDS = 30

# ============================================================
# Data Models
# ============================================================

class JobD(BaseModel):
    role: str | None = None 
    required_skills: list[str] = []
    preferred_skills: list[str] = []
    minimum_experience: float | None = None 
    education_requirements: list[str] = []
    responsibilities: list[str] = []


class Experience(BaseModel):
    company: str | None = None 
    role: str | None = None 
    duration: str | None = None 
    description: str | None = None 
    skills_used: list[str] = []


class Resume(BaseModel):
    name: str | None = None 
    email: str | None = None 
    phone: str | None = None 
    total_experience_years: float | None = None 
    skills: list[str] = []
    experiences: list[Experience] = []
    education: list[str] = []
    projects: list[str] = []
    certifications: list[str] = []


class MatchResult(BaseModel):
    matching_skills: list[str] = []
    missing_skills: list[str] = []
    experience_met: bool | None = None 
    score: float = 0.0
    verdict: str = "" 


jobd_schema = JobD.model_json_schema()
resume_schema = Resume.model_json_schema()
match_schema = MatchResult.model_json_schema()


# ============================================================
# LLM Calls 
# ============================================================


def call_groq(system_prompt, user_prompt):
    messages = []
    if system_prompt:
        messages.append({"role": "system", "content": system_prompt})
    messages.append({"role": "user", "content": user_prompt})

    response = client.chat.completions.create(
        model = model,
        messages = messages, 
        response_format = {"type": "json_object"},
    )
    return json.loads(response.choices[0].message.content)


def parse_jd(jd_text):
    system_prompt = f"""
You are an expert HR assistant.
Analyse the job description and extract structured information from it.
Return ONLY valid JSON matching this schema.
{jobd_schema}
 
IMPORTANT:
Do NOT return the schema itself.
Do NOT return fields like "properties", "title" or "type".
Fill the schema with actual information extracted from the job description.
If minimum experience is not mentioned, return null.
If information for a list is missing, return an empty list.
Do not invent information.
"""

    user_prompt = f"Analyse the following job description:\n\n{jd_text}"
    data = call_groq(system_prompt, user_prompt)
    return JobD(**data)


def parse_resume(resume_text):
    system_prompt = f"""
You are an expert resume parser.
Extract information from the resume based on its meaning, not only exact section headings.
Different resumes may use different headings, e.g. Experience, Professional Experience,
Work History, Employment, Internships. Skills may appear in the skills section, work
experience, internships, or projects.
 
Return ONLY valid JSON matching this schema:
{resume_schema}
 
Important rules:
1. Do not invent information.
2. If a value is not available, return null.
3. If a list has no information, return an empty list.
4. Include internships inside experiences.
5. Extract skills mentioned across the entire resume.
"""

    user_prompt = f"Parse the following resume:\n\n{resume_text}"
    data = call_groq(system_prompt, user_prompt)
    return Resume(**data)


def match_resume_to_job(job, resume):
    system_prompt = f"""
You are an HR recruiter comparing a candidate's resume to a job description.
Return ONLY valid JSON matching this schema:
{match_schema}
 
Rules:
1. matching_skills: skills present in both the resume and the job's required/preferred skills.
2. missing_skills: important required skills from the job that are NOT in the resume.
3. experience_met: true/false, or null if minimum_experience was not stated.
4. score: overall match percentage from 0 to 100. Weight required skills more heavily
   than preferred skills.
5. verdict: one short sentence summarising overall fit.
Do not invent information.
"""

    user_prompt = f"""
JOB DESCRIPTION:
{job.model_dump_json(indent=2)}
 
CANDIDATE RESUME:
{resume.model_dump_json(indent=2)}
"""

    data = call_groq(system_prompt, user_prompt)
    return MatchResult(**data)


# ============================================================
# Usage Limiting 
# ============================================================

def check_cooldown():
    last = st.session_state.get("last_request_time")
    if last:
        elapsed = (datetime.now() - last).total_seconds()
        if elapsed < COOLDOWN_SECONDS:
            wait = int(COOLDOWN_SECONDS - elapsed)
            st.warning(f"Please wait {wait} more second(s) before analysing again.")
            return False 
    return True


def check_length(text, label):
    if len(text) > MAX_CHARS:
        st.error(f"{label} is too long ({len(text)} characters). Please keep it under {MAX_CHARS} characters.")
        return False 
    return True 


# ============================================================
# File Reading (PDF/DOCX)
# ============================================================

def read_pdf(uploaded_file):
    reader = PdfReader(uploaded_file)
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"
    return text 


def read_docx(uploaded_file):
    document = Document(uploaded_file)
    text = ""
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"
    return text 


def extract_text_from_upload(uploaded_file):
    name = uploaded_file.name.lower()
    if name.endswith(".pdf"):
        return read_pdf(uploaded_file)
    elif name.endswith(".docx"):
        return read_docx(uploaded_file)
    return None 


# ============================================================
# Visual Helpers
# ============================================================

INK = "#101318"
INK_SOFT = "#191D25"
PAPER = "#F7F2E8"
PAPER_DARK = "#EEE5D4"
INK_TEXT = "#29261F"

GOLD = "#D2A64A"
GOLD_DARK = "#B8892F"
GREEN = "#6F9A7A"
RED = "#C96B61"

MUTED = "#AAA69C"


# ============================================================
# Page Configuration
# ============================================================

st.set_page_config(
    page_title="Resume ↔ Job Matcher",
    page_icon="🧩",
    layout="wide",
)


# ============================================================
# CSS
# ============================================================

CUSTOM_CSS = f"""
<style>

/* ------------------------------------------------------------
   GLOBAL
------------------------------------------------------------ */

.stApp {{
    background:
        radial-gradient(
            circle at 10% 0%,
            rgba(210, 166, 74, 0.055),
            transparent 30%
        ),
        radial-gradient(
            circle at 90% 100%,
            rgba(111, 154, 122, 0.045),
            transparent 30%
        ),
        {INK};
}}

.block-container {{
    max-width: 1150px;
    padding-top: 3rem;
    padding-bottom: 4rem;
}}


/* ------------------------------------------------------------
   TYPOGRAPHY
------------------------------------------------------------ */

h1 {{
    font-family: Georgia, 'Times New Roman', serif !important;
    color: #F5F0E7 !important;
    font-size: 3.1rem !important;
    font-weight: 600 !important;
    letter-spacing: -0.045em !important;
    line-height: 1.08 !important;
    margin-bottom: 0.4rem !important;
}}

h2, h3 {{
    font-family: Georgia, 'Times New Roman', serif !important;
    color: #F5F0E7 !important;
}}

p {{
    color: {MUTED};
}}

.stCaption {{
    color: {MUTED} !important;
}}


/* ------------------------------------------------------------
   TEXT AREAS
------------------------------------------------------------ */

textarea {{
    background-color: {PAPER} !important;
    color: {INK_TEXT} !important;

    border: 1px solid #D8CEBB !important;
    border-radius: 14px !important;

    font-family: Arial, sans-serif !important;
    font-size: 0.94rem !important;
    line-height: 1.6 !important;

    padding: 16px !important;

    box-shadow:
        0 3px 12px rgba(0, 0, 0, 0.08) !important;
}}

textarea::placeholder {{
    color: #817A70 !important;
    opacity: 1 !important;
}}

textarea:focus {{
    border: 1px solid {GOLD} !important;

    box-shadow:
        0 0 0 3px rgba(210, 166, 74, 0.13),
        0 5px 18px rgba(0, 0, 0, 0.12) !important;
}}


/* ------------------------------------------------------------
   RADIO
------------------------------------------------------------ */

[data-testid="stRadio"] {{
    margin-bottom: 0.5rem;
}}

[data-testid="stRadio"] label {{
    color: #C9C4B8 !important;
    font-size: 0.88rem !important;
}}


/* ------------------------------------------------------------
   FILE UPLOADER
------------------------------------------------------------ */

[data-testid="stFileUploaderDropzone"] {{
    background: {PAPER} !important;
    border: 1.5px dashed #C8B994 !important;
    border-radius: 14px !important;
    min-height: 180px !important;
    padding: 1.5rem !important;
}}

/* ------------------------------------------------------------
   UPLOAD BUTTON
------------------------------------------------------------ */

[data-testid="stFileUploaderDropzone"] button {{
    background: {INK_SOFT} !important;
    color: #F7F2E8 !important;
    border: 1px solid rgba(210, 166, 74, 0.45) !important;
    border-radius: 10px !important;
    padding: 0.55rem 1.15rem !important;
    font-size: 0.88rem !important;
    font-weight: 600 !important;
}}

[data-testid="stFileUploaderDropzone"] button * {{
    color: #F7F2E8 !important;
}}

[data-testid="stFileUploaderDropzone"] button svg {{
    color: {GOLD} !important;
}}

[data-testid="stFileUploaderDropzone"] button:hover {{
    background: #222731 !important;
    border-color: {GOLD} !important;
}}


/* ------------------------------------------------------------
   UPLOADER INSTRUCTION TEXT
------------------------------------------------------------ */

/* This is the important part */
[data-testid="stFileUploaderDropzoneInstructions"] {{
    color: #6F685D !important;
    opacity: 1 !important;
}}

[data-testid="stFileUploaderDropzoneInstructions"] * {{
    color: #6F685D !important;
    opacity: 1 !important;
}}

/* Streamlit sometimes renders the size/type text as <small> */
[data-testid="stFileUploaderDropzoneInstructions"] small {{
    color: #6F685D !important;
    opacity: 1 !important;
}}

/* And sometimes as paragraph text */
[data-testid="stFileUploaderDropzoneInstructions"] p {{
    color: #6F685D !important;
    opacity: 1 !important;
}}


/* ------------------------------------------------------------
   FALLBACK — ANY NON-BUTTON TEXT IN DROPZONE
------------------------------------------------------------ */

[data-testid="stFileUploaderDropzone"] > section > div:not(:has(button)) {{
    color: #6F685D !important;
}}

[data-testid="stFileUploaderDropzone"] > section > div:not(:has(button)) * {{
    color: #6F685D !important;
    opacity: 1 !important;
}}


/* ------------------------------------------------------------
   PRIMARY BUTTON
------------------------------------------------------------ */

button[kind="primary"] {{
    background: linear-gradient(
        135deg,
        {GOLD},
        {GOLD_DARK}
    ) !important;

    color: #17140F !important;

    border: none !important;
    border-radius: 12px !important;

    min-height: 50px !important;

    font-size: 0.96rem !important;
    font-weight: 700 !important;

    box-shadow:
        0 8px 22px rgba(210, 166, 74, 0.20) !important;

    transition: all 0.15s ease !important;
}}

button[kind="primary"] p {{
    color: #17140F !important;
}}

button[kind="primary"]:hover {{
    filter: brightness(1.07) !important;
    transform: translateY(-2px) !important;

    box-shadow:
        0 11px 28px rgba(210, 166, 74, 0.28) !important;
}}


/* ------------------------------------------------------------
   CONTAINERS
------------------------------------------------------------ */

[data-testid="stVerticalBlockBorderWrapper"] {{
    background: {INK_SOFT};
    border: 1px solid rgba(255, 255, 255, 0.065);
    border-radius: 16px;
    padding: 1.15rem;
}}


/* ------------------------------------------------------------
   ALERTS
------------------------------------------------------------ */

[data-testid="stAlert"] {{
    border-radius: 10px !important;
}}


/* ------------------------------------------------------------
   HIDE STREAMLIT CHROME
------------------------------------------------------------ */

#MainMenu {{
    visibility: hidden;
}}

footer {{
    visibility: hidden;
}}

header {{
    background: transparent !important;
}}

</style>
"""

st.markdown(CUSTOM_CSS, unsafe_allow_html=True)


# ============================================================
# HEADER
# ============================================================

st.markdown("### ✦ CAREER TOOL")

st.title("Resume ↔ Job Description")

st.markdown(
    """
    See how closely your resume aligns with a role —
    from matching skills to the gaps worth working on.
    """
)

st.write("")


# ============================================================
# INPUT AREA
# ============================================================

col1, col2 = st.columns(2, gap="large")


# ============================================================
# JOB DESCRIPTION
# ============================================================

with col1:

    st.subheader("💼 Job description")

    st.caption("What is the company looking for?")

    jd_text = st.text_area(
        "Job description",
        height=300,
        placeholder=(
            "Paste the job description here...\n\n"
            "For example:\n"
            "• Python\n"
            "• SQL\n"
            "• Machine Learning\n"
            "• Data Analysis\n"
            "• 1–2 years experience"
        ),
        label_visibility="collapsed",
    )


# ============================================================
# RESUME
# ============================================================

with col2:

    st.subheader("📄 Your resume")

    st.caption("What are you bringing to the table?")

    resume_method = st.radio(
        "Resume input method",
        ["Paste text", "Upload PDF/DOCX"],
        horizontal=True,
        label_visibility="collapsed",
    )

    resume_text = ""

    if resume_method == "Paste text":

        resume_text = st.text_area(
            "Resume text",
            height=300,
            placeholder=(
                "Paste your resume here...\n\n"
                "Tip: include your skills, experience, "
                "projects and education."
            ),
            label_visibility="collapsed",
        )

    else:

        uploaded_resume = st.file_uploader(
            "Upload resume",
            type=["pdf", "docx"],
            label_visibility="collapsed",
        )

        if uploaded_resume is not None:

            extracted = extract_text_from_upload(
                uploaded_resume
            )

            if extracted:

                resume_text = extracted

                st.success(
                    f"✓ Extracted {len(resume_text)} characters "
                    f"from {uploaded_resume.name}"
                )

            else:

                st.error(
                    "Could not read that file. "
                    "Please upload a PDF or DOCX."
                )


# ============================================================
# CTA
# ============================================================

st.write("")

st.caption("Ready to see where you stand?")

button_left, button_middle, button_right = st.columns(
    [1, 1.2, 1]
)

with button_middle:

    analyse = st.button(
        "✦  Analyse my match",
        type="primary",
        use_container_width=True,
    )


# ============================================================
# ANALYSIS
# ============================================================

if analyse:

    # --------------------------------------------------------
    # Validation
    # --------------------------------------------------------

    if not jd_text.strip() or not resume_text.strip():

        st.error(
            "Please paste both a job description and a resume."
        )

    elif (
        not check_length(
            jd_text,
            "Job description",
        )
        or not check_length(
            resume_text,
            "Resume",
        )
    ):

        pass

    elif not check_cooldown():

        pass

    else:

        st.session_state["last_request_time"] = datetime.now()

        try:

            # ------------------------------------------------
            # Parse JD
            # ------------------------------------------------

            with st.spinner(
                "Parsing job description..."
            ):

                job = parse_jd(jd_text)


            # ------------------------------------------------
            # Parse Resume
            # ------------------------------------------------

            with st.spinner(
                "Parsing resume..."
            ):

                resume = parse_resume(resume_text)


            # ------------------------------------------------
            # Compare
            # ------------------------------------------------

            with st.spinner(
                "Comparing your resume with the role..."
            ):

                result = match_resume_to_job(
                    job,
                    resume,
                )


        except Exception as e:

            msg = str(e).lower()

            if "rate limit" in msg or "429" in msg:

                st.error(
                    "The analysis service is at capacity right now "
                    "(rate limit reached). Please try again in a minute."
                )

            else:

                st.error(
                    f"Something went wrong while analysing: {e}"
                )


        else:

            # =================================================
            # SCORE
            # =================================================

            st.write("")

            if result.score >= 75:

                score_icon = "✦"
                score_label = "Strong match"
                score_color = GREEN

            elif result.score >= 50:

                score_icon = "◐"
                score_label = "Moderate match"
                score_color = GOLD

            else:

                score_icon = "↗"
                score_label = "Needs work"
                score_color = RED


            score_col1, score_col2 = st.columns(
                [1, 4],
                gap="large",
            )


            with score_col1:

                st.metric(
                    "MATCH",
                    f"{result.score:.0f}%",
                )


            with score_col2:

                st.subheader(
                    f"{score_icon} {score_label}"
                )

                st.write(result.verdict)


            st.divider()


            # =================================================
            # SKILLS
            # =================================================

            left_result, right_result = st.columns(
                2,
                gap="large",
            )


            # -------------------------------------------------
            # MATCHING SKILLS
            # -------------------------------------------------

            with left_result:

                with st.container(border=True):

                    st.subheader(
                        "✓ Matching skills"
                    )

                    if result.matching_skills:

                        for skill in result.matching_skills:

                            st.markdown(
                                f"• **{skill}**"
                            )

                    else:

                        st.caption(
                            "No matching skills found."
                        )


                # -------------------------------------------------
                # RESUME SKILLS
                # -------------------------------------------------

                with st.container(border=True):

                    st.subheader(
                        "✧ Skills in your resume"
                    )

                    if resume.skills:

                        for skill in resume.skills:

                            st.markdown(
                                f"• {skill}"
                            )

                    else:

                        st.caption(
                            "No skills were identified."
                        )


            # -------------------------------------------------
            # MISSING / REQUIRED
            # -------------------------------------------------

            with right_result:

                with st.container(border=True):

                    st.subheader(
                        "△ Missing skills"
                    )

                    if result.missing_skills:

                        for skill in result.missing_skills:

                            st.markdown(
                                f"• **{skill}**"
                            )

                    else:

                        st.caption(
                            "No obvious skill gaps found."
                        )


                # -------------------------------------------------
                # REQUIRED SKILLS
                # -------------------------------------------------

                with st.container(border=True):

                    st.subheader(
                        "◎ Required skills"
                    )

                    if job.required_skills:

                        for skill in job.required_skills:

                            st.markdown(
                                f"• {skill}"
                            )

                    else:

                        st.caption(
                            "No required skills were identified."
                        )


            # =================================================
            # EXPERIENCE
            # =================================================

            if job.minimum_experience is not None:

                st.write("")

                if result.experience_met is True:

                    experience_text = "✓ Experience requirement met"

                elif result.experience_met is False:

                    experience_text = (
                        "△ Experience requirement not clearly met"
                    )

                else:

                    experience_text = (
                        "◐ Experience requirement is unclear"
                    )


                with st.container(border=True):

                    st.subheader(
                        "Experience"
                    )

                    st.write(
                        f"Required: **{job.minimum_experience} years**"
                    )

                    st.caption(
                        experience_text
                    )
